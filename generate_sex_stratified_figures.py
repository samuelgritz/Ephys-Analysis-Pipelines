"""
generate_sex_stratified_figures.py
===================================
Generates sex-stratified (Male vs Female) figures and statistics for the GNB1 manuscript.
For each metric, plots are broken down by Genotype × Sex (4 groups: WT-M, WT-F, I80T/+-M, I80T/+-F).
Aesthetics match the exact publication style of the main paper (6.93 in / 3.45 in dimensions,
Arial typography, clean spines, outwards ticks, anchored clean Y-ticks, and Okabe-Ito colors).

Metrics covered:
  1. Body Weights (P8-P10, P28, Adult)
  2. Summed Activity & Hourly Activity – DVC Circadian
  3. FI Midpoint (Somatic Excitability)
  4. Plateau Area (Dendritic Excitability)
  5. Open Field Locomotion (Distance) & Open Field Anxiety (Center:Outer Ratio)
  6. T-Maze Behavior (Distance, Total Arm Entries, % Alternation)
  7. E/I Imbalance – All three pathways (Perforant, Schaffer, Basal_Stratum_Oriens)
  8. OLM (Object Location Memory) – Testing Discrimination Index

Stats (Mann-Whitney U) are computed for:
  - Male:   WT-M vs I80T/+-M
  - Female: WT-F vs I80T/+-F
Results are appended to paper_data/Master_Stats_Summary.csv and .xlsx as new tabs.

Figures are saved to paper_figures/ as .png, .svg, and .pdf.
"""

import os
import sys
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
from scipy import stats

# Import plotting utilities (same style as generate_figures.py)
from plotting_utils import (
    setup_publication_style,
    save_current_fig,
    rename_genotype,
    apply_clean_yticks,
    get_safe_y,
    add_subplot_label,
    draw_significance,
    COLORS,
)

# ─────────────────────────────────────────────────────────────────────────────
# CONFIG
# ─────────────────────────────────────────────────────────────────────────────
PAPER_DATA_DIR  = "paper_data"
OUTPUT_FIG_DIR  = "paper_figures"
MASTER_CSV      = os.path.join(PAPER_DATA_DIR, "Master_Stats_Summary.csv")
MASTER_XLSX     = os.path.join(PAPER_DATA_DIR, "Master_Stats_Summary.xlsx")

# Genotype × Sex groups
SEX_GENO_ORDER  = ["WT-M", "WT-F", "I80T/+-M", "I80T/+-F"]

# Okabe-Ito colorblind-safe palette (retained for sex stratification)
# WT   → blue family:    sky-blue (Male lighter) / dark-blue (Female darker)
# I80T → orange family:  orange   (Male lighter) / vermilion (Female darker)
SEX_GENO_COLORS = {
    "WT-M":     "#56B4E9",   # sky blue
    "WT-F":     "#0072B2",   # dark blue
    "I80T/+-M": "#E69F00",   # orange
    "I80T/+-F": "#D55E00",   # vermilion
}
SEX_GENO_LINES = {           # linestyle by sex (reinforces color)
    "WT-M":     "--",
    "WT-F":     "-",
    "I80T/+-M": "--",
    "I80T/+-F": "-",
}


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def p_to_sig(p):
    try:
        p = float(p)
    except (ValueError, TypeError):
        return "?"
    if p < 0.001:   return "***"
    elif p < 0.01:  return "**"
    elif p < 0.05:  return "*"
    else:           return "ns"


def mann_whitney(a, b):
    """Returns (U, p). Returns (nan, nan) if either group is too small."""
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    a = a[~np.isnan(a)]
    b = b[~np.isnan(b)]
    if len(a) < 2 or len(b) < 2:
        return np.nan, np.nan
    u, p = stats.mannwhitneyu(a, b, alternative="two-sided")
    return u, p


def make_sex_geno_col(df, genotype_col="Genotype", sex_col="Sex"):
    """Creates a combined Genotype-Sex label column for grouping."""
    df = df.copy()
    sex_map = {"M": "M", "Male": "M", "male": "M",
               "F": "F", "Female": "F", "female": "F"}
    df["_Sex"]  = df[sex_col].astype(str).str.strip().map(sex_map)
    df["_Geno"] = df[genotype_col]
    df["SexGeno"] = df["_Geno"] + "-" + df["_Sex"]
    return df


def plot_sex_bar_scatter(ax, data, y_col, title, ylabel,
                         order=None, show_pval_bracket=True):
    """
    Bar + scatter plot for 4 sex×genotype groups matching exact paper formatting.
    Scatter points aligned at fixed x-coordinates (no random jitter).
    Returns a list of stats dicts: [{group_comparison, U, p, sig}, ...].
    """
    if order is None:
        order = SEX_GENO_ORDER
    color_map = SEX_GENO_COLORS

    bar_width = 0.55
    stats_out = []

    # Marker shape encodes sex: Male = triangle (^), Female = circle (o)
    MARKER_MAP = {"WT-M": "^", "WT-F": "o", "I80T/+-M": "^", "I80T/+-F": "o"}

    for i, group in enumerate(order):
        subset = data[data["SexGeno"] == group]
        if subset.empty:
            continue
        values = subset[y_col].dropna().values
        if len(values) == 0:
            continue
        color  = color_map.get(group, "gray")
        marker = MARKER_MAP.get(group, "o")
        mean   = np.mean(values)
        sem    = np.std(values, ddof=1) / np.sqrt(len(values)) if len(values) > 1 else 0

        # Bar matching paper style: alpha=0.5, no border
        ax.bar(i, mean, width=bar_width, color=color, alpha=0.50, edgecolor="none")
        # Errorbar matching paper style: capsize=1.5, elinewidth=0.8
        ax.errorbar(i, mean, yerr=sem, fmt="o", color=color,
                    capsize=1.5, capthick=0.8, elinewidth=0.8, markersize=2)
        # Scatter points aligned at fixed x-coordinate (no jitter)
        ax.scatter(np.full(len(values), i), values,
                   color=color, marker=marker, s=9, zorder=3, alpha=0.75,
                   linewidths=0.3, edgecolors="white")

    ax.set_xticks(range(len(order)))
    labels = []
    for g in order:
        sub = data[data["SexGeno"] == g].dropna(subset=[y_col])
        n   = len(sub)
        labels.append(f"{g}\n(n={n})")
    ax.set_xticklabels(labels, fontsize=7)
    ax.set_ylabel(ylabel, fontsize=8)
    ax.set_title(title, fontsize=9, fontweight="bold" if len(title) < 25 else "normal")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(0.8)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.tick_params(direction="out", length=3, width=0.8, labelsize=8)
    ax.grid(False)

    # Smart Y limits and clean ticks
    all_vals = data[y_col].dropna().values
    if len(all_vals):
        lo, hi = all_vals.min(), all_vals.max()
        rng = hi - lo if hi != lo else abs(hi) or 1
        if lo >= 0:
            ax.set_ylim(0, hi + rng * 0.25)
        elif hi <= 0:
            ax.set_ylim(lo - rng * 0.25, 0)
        else:
            ax.set_ylim(lo - rng * 0.2, hi + rng * 0.25)
    apply_clean_yticks(ax)

    # ── Significance brackets ──────────────────────────────────────────────
    if show_pval_bracket:
        ylo, yhi = ax.get_ylim()
        y_range  = yhi - ylo

        # Male comparison: WT-M (idx=0) vs I80T/+-M (idx=2)
        wt_m  = data[data["SexGeno"] == "WT-M"][y_col].dropna().values
        mut_m = data[data["SexGeno"] == "I80T/+-M"][y_col].dropna().values
        u_m, p_m = mann_whitney(wt_m, mut_m)
        if not np.isnan(p_m):
            y_br_m = yhi - y_range * 0.05
            draw_significance(ax, 0, 2, p_m, y_br_m, bracket=True)
            stats_out.append(dict(Comparison="WT-M vs I80T/+-M", U=u_m, P=p_m, Sig=p_to_sig(p_m)))

        # Female comparison: WT-F (idx=1) vs I80T/+-F (idx=3)
        wt_f  = data[data["SexGeno"] == "WT-F"][y_col].dropna().values
        mut_f = data[data["SexGeno"] == "I80T/+-F"][y_col].dropna().values
        u_f, p_f = mann_whitney(wt_f, mut_f)
        if not np.isnan(p_f):
            y_br_f = yhi - y_range * 0.12
            draw_significance(ax, 1, 3, p_f, y_br_f, bracket=True)
            stats_out.append(dict(Comparison="WT-F vs I80T/+-F", U=u_f, P=p_f, Sig=p_to_sig(p_f)))

    return stats_out


def save_fig_png_svg(fig, fig_name):
    """Saves figure in .pdf, .svg, and .png format using publication exporter."""
    save_current_fig(fig_name)


def build_stats_row(figure, subpanel, metric, pathway, condition,
                    data, y_col, sex_label, comparison):
    """Build a stats dict in Master_Stats_Summary format for one sex comparison."""
    geno_a, geno_b = comparison.split(" vs ")
    grp_a = data[data["SexGeno"] == geno_a][y_col].dropna()
    grp_b = data[data["SexGeno"] == geno_b][y_col].dropna()
    u, p  = mann_whitney(grp_a.values, grp_b.values)

    def _f(x): return round(float(x), 4) if pd.notna(x) and not np.isnan(float(x)) else np.nan

    return dict(
        Figure=figure, Subpanel=subpanel, Metric=metric + f" [{sex_label}]",
        Pathway=pathway, Condition=condition,
        WT_Mean=_f(grp_a.mean()), WT_SEM=_f(grp_a.sem()),
        WT_N=len(grp_a),
        I80T_Mean=_f(grp_b.mean()), I80T_SEM=_f(grp_b.sem()),
        I80T_N=len(grp_b),
        Test_Used="Mann-Whitney U (sex-stratified)",
        Statistic=_f(u),
        P_Value=_f(p),
        Significance=p_to_sig(p),
        Notes=f"Sex-stratified: {comparison}",
    )


# ─────────────────────────────────────────────────────────────────────────────
# DATA LOADING
# ─────────────────────────────────────────────────────────────────────────────

def load_all_data():
    """Load all required datasets and attach sex info where needed."""
    data = {}

    # master_df (contains Cell_ID and Sex for physiology cells)
    master = pd.read_csv("master_df.csv", low_memory=False)
    master["Cell_ID"] = master["Cell_ID"].astype(str).str.strip()
    master["Sex"]     = master["Sex"].astype(str).str.strip()
    data["master"] = master

    # 1. Weights
    df_w = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                    "Mouse_Weights_Processed.csv"))
    df_w = rename_genotype(df_w)
    df_w = make_sex_geno_col(df_w, sex_col="Sex")
    data["weights"] = df_w

    # 2. DVC – Summed Activity & Hourly Activity
    df_cage = pd.read_csv(os.path.join(PAPER_DATA_DIR, "DVC_Analysis",
                                        "Cage_Specific_Hours_Summary.csv"))
    df_cage = rename_genotype(df_cage)
    df_cage = make_sex_geno_col(df_cage, sex_col="Sex")
    data["dvc_cage"] = df_cage

    df_dvc_hourly = pd.read_csv(os.path.join(PAPER_DATA_DIR, "DVC_Analysis",
                                             "Hourly_Stats_By_Sex.csv"))
    df_dvc_hourly = rename_genotype(df_dvc_hourly)
    df_dvc_hourly["Sex_Clean"] = df_dvc_hourly["Sex"].astype(str).str.strip().map({"M": "M", "Male": "M", "F": "F", "Female": "F"})
    df_dvc_hourly["SexGeno"]   = df_dvc_hourly["Genotype"] + "-" + df_dvc_hourly["Sex_Clean"]
    data["dvc_hourly"] = df_dvc_hourly

    # 3. FI Midpoint – join sex from master_df
    df_fi = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Firing_Rate",
                                      "Sigmoid_Fit_Params.csv"))
    df_fi["Cell_ID"] = df_fi["Cell_ID"].astype(str).str.strip()
    sex_map = master[["Cell_ID", "Sex"]].drop_duplicates()
    df_fi = df_fi.merge(sex_map, on="Cell_ID", how="left")
    df_fi = df_fi.rename(columns={"Midpoint": "FI_Midpoint"}) \
                  if "FI_Midpoint" not in df_fi.columns else df_fi
    df_fi = rename_genotype(df_fi)
    df_fi = make_sex_geno_col(df_fi, sex_col="Sex")
    data["fi_midpoint"] = df_fi

    # 4. E/I Imbalance – already has Sex col in E_I_amplitudes.csv
    df_ei = pd.read_csv(os.path.join(PAPER_DATA_DIR, "E_I_data",
                                      "E_I_amplitudes.csv"))
    df_ei = rename_genotype(df_ei)
    df_ei = make_sex_geno_col(df_ei, sex_col="Sex")
    data["ei"] = df_ei

    # 5. Plateau Area
    df_plat = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Plateau_data",
                                         "Plateau_data.csv"))
    df_plat = rename_genotype(df_plat)
    if "Sex" in df_plat.columns:
        df_plat = make_sex_geno_col(df_plat, sex_col="Sex")
    data["plateau"] = df_plat

    # 6. Open Field Locomotion & Anxiety
    df_of_loc = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                           "Open_Field_Locomotion_Trial1.csv"))
    df_of_anx = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                           "Open_Field_Anxiety_Processed.csv"))
    df_of_loc = rename_genotype(df_of_loc)
    df_of_anx = rename_genotype(df_of_anx)
    df_of_loc = make_sex_geno_col(df_of_loc, sex_col="Sex")
    df_of_anx = make_sex_geno_col(df_of_anx, sex_col="Sex")
    data["of_loc"] = df_of_loc
    data["of_anx"] = df_of_anx

    # 7. T-Maze
    df_tmaze_alt = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                            "T_Maze_Alternations.csv"))
    df_tmaze_ent = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                            "T_Maze_Zone_Entries.csv"))
    df_tmaze_alt = rename_genotype(df_tmaze_alt)
    df_tmaze_ent = rename_genotype(df_tmaze_ent)
    df_tmaze_alt = make_sex_geno_col(df_tmaze_alt, sex_col="Sex")
    df_tmaze_ent = make_sex_geno_col(df_tmaze_ent, sex_col="Sex")
    df_tmaze_ent["Total_Arm_Entries"] = (df_tmaze_ent["Start : entries"] + 
                                         df_tmaze_ent["Left Arm : entries"] + 
                                         df_tmaze_ent["Right Arm : entries"])
    data["tmaze_alt"] = df_tmaze_alt
    data["tmaze_ent"] = df_tmaze_ent

    # 8. OLM
    df_olm = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Behavior_Analysis",
                                        "OLM_Summary_Deltas.csv"))
    df_olm = rename_genotype(df_olm)
    df_olm = make_sex_geno_col(df_olm, sex_col="Sex")
    data["olm"] = df_olm

    # 9. Somatic Intrinsic Properties (Input Resistance)
    df_intr = pd.read_csv(os.path.join(PAPER_DATA_DIR, "Physiology_Analysis",
                                       "Intrinsic_properties.csv"))
    df_intr = rename_genotype(df_intr)
    df_intr = make_sex_geno_col(df_intr, sex_col="Sex")
    data["intrinsic"] = df_intr

    return data


# ─────────────────────────────────────────────────────────────────────────────
# FIGURE GENERATORS
# ─────────────────────────────────────────────────────────────────────────────

def plot_sex_longitudinal_weights_single(ax, df_w, sex_code, sex_title):
    """Plot longitudinal developmental body weight line plot (P8-P10, P28, Adult) for Male or Female matching Fig 1B."""
    time_order = ["P8-P10", "P28", "Adult"]
    x_labels = ["P8-P10", "P28", "Adult (P49)"]
    time_map = {label: i for i, label in enumerate(time_order)}

    df_sex = df_w[df_w["Sex"].astype(str).str.strip().isin([sex_code, f"{sex_code}ale"])].copy()

    # Calculate summary mean and SEM
    summary = df_sex.groupby(["Timepoint_Label", "Genotype"])["Weight_g"].agg(["mean", "sem", "count"]).reset_index()
    summary["x_pos"] = summary["Timepoint_Label"].map(time_map)

    geno_order = ["WT", "I80T/+"]
    max_h = 0
    stats_out = []

    for group in geno_order:
        sub = summary[summary["Genotype"] == group].sort_values("x_pos")
        if sub.empty:
            continue
        key = f"{group}-{sex_code}"
        color = SEX_GENO_COLORS.get(key, "black" if group == "WT" else "red")
        
        x = sub["x_pos"].to_numpy()
        y = sub["mean"].to_numpy()
        e = sub["sem"].to_numpy()
        
        n_min, n_max = sub["count"].min(), sub["count"].max()
        n_str = f"{n_min}–{n_max}" if n_min != n_max else str(n_min)
        
        ax.errorbar(x, y, yerr=e, fmt='o', color=color, capsize=2, capthick=0.8,
                    markersize=4, elinewidth=0.8, label=f"{group} (n={n_str})")
        ax.plot(x, y, color=color, linestyle='-', linewidth=1.2)
        
        curr_max = (sub["mean"] + sub["sem"]).max()
        if curr_max > max_h:
            max_h = curr_max

    # Annotate significance between WT and I80T/+ at each timepoint
    for i, tp in enumerate(time_order):
        sub_tp = df_sex[df_sex["Timepoint_Label"] == tp]
        wt_vals = sub_tp[sub_tp["Genotype"] == "WT"]["Weight_g"].dropna().values
        mut_vals = sub_tp[sub_tp["Genotype"] == "I80T/+"]["Weight_g"].dropna().values
        if len(wt_vals) > 0 and len(mut_vals) > 0:
            u, p = mann_whitney(wt_vals, mut_vals)
            sig = p_to_sig(p)
            
            tp_sum = summary[summary["Timepoint_Label"] == tp]
            top_y = (tp_sum["mean"] + tp_sum["sem"]).max() if not tp_sum.empty else 25.0
            
            if sig != "ns":
                ax.text(i, top_y + 1.0, sig, ha='center', va='bottom', fontsize=9, fontweight='bold')
            else:
                ax.text(i, top_y + 1.0, "ns", ha='center', va='bottom', fontsize=7, color='gray')
            
            stats_out.append(dict(Timepoint=tp, Sex=sex_title, U=u, P=p, Sig=sig))

    ax.set_xticks(range(len(time_order)))
    ax.set_xticklabels(x_labels, fontsize=8)
    ax.set_ylabel("Weight (g)", fontsize=8)
    ax.set_title(f"Developmental Body Weight ({sex_title})", fontsize=9, fontweight="bold")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(0.8)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.tick_params(direction="out", length=3, width=0.8, labelsize=8)
    ax.set_ylim(0, max(28, max_h + 4.5))
    ax.legend(frameon=False, fontsize=7, loc="upper left")
    return stats_out


def plot_sex_weights_line_plots(data):
    """Figure: Developmental Body Weights (Longitudinal Line Plots only)."""
    df_w = data["weights"].dropna(subset=["Weight_g", "SexGeno"])
    fig, axes = plt.subplots(1, 2, figsize=(6.93, 3.2))
    fig.subplots_adjust(wspace=0.35, left=0.10, right=0.96, top=0.88, bottom=0.22)
    st_f = plot_sex_longitudinal_weights_single(axes[0], df_w, "F", "Females")
    st_m = plot_sex_longitudinal_weights_single(axes[1], df_w, "M", "Males")
    add_subplot_label(axes[0], "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(axes[1], "B", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    return fig, st_f + st_m


def plot_sex_weights_bar_plots(data):
    """Figure: Developmental Body Weights (Timepoint Bar Plots only)."""
    df_w = data["weights"].dropna(subset=["Weight_g", "SexGeno"])
    timepoints = ["P8-P10", "P28", "Adult"]
    fig, axes = plt.subplots(1, 3, figsize=(6.93, 2.8))
    fig.subplots_adjust(wspace=0.35, left=0.09, right=0.96, top=0.88, bottom=0.22)
    stats_all = []
    for i, tp in enumerate(timepoints):
        sub = df_w[df_w["Timepoint_Label"] == tp]
        st = plot_sex_bar_scatter(axes[i], sub, "Weight_g", f"Weight ({tp})", "Weight (g)")
        stats_all.extend(st)
        add_subplot_label(axes[i], chr(65 + i), x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    return fig, stats_all


def plot_sex_weights(data):
    """Figure: Developmental Body Weights combining BOTH Longitudinal Line Plots and Timepoint Bar Plots."""
    print("  → Body Weight by Sex (Line Plots + Bar Plots)")
    df_w = data["weights"].dropna(subset=["Weight_g", "SexGeno"])

    fig = plt.figure(figsize=(6.93, 5.8))
    gs  = gridspec.GridSpec(2, 1, height_ratios=[1.0, 1.0], hspace=0.48)

    # Row 1: Longitudinal Line Plots (Females, Males)
    gs_r1 = gridspec.GridSpecFromSubplotSpec(1, 2, subplot_spec=gs[0], wspace=0.35)
    ax_wf = fig.add_subplot(gs_r1[0])
    ax_wm = fig.add_subplot(gs_r1[1])
    st_f  = plot_sex_longitudinal_weights_single(ax_wf, df_w, "F", "Females")
    st_m  = plot_sex_longitudinal_weights_single(ax_wm, df_w, "M", "Males")

    # Row 2: Timepoint Bar Plots (P8-P10, P28, Adult)
    gs_r2 = gridspec.GridSpecFromSubplotSpec(1, 3, subplot_spec=gs[1], wspace=0.35)
    ax_b1 = fig.add_subplot(gs_r2[0])
    ax_b2 = fig.add_subplot(gs_r2[1])
    ax_b3 = fig.add_subplot(gs_r2[2])

    st1 = plot_sex_bar_scatter(ax_b1, df_w[df_w["Timepoint_Label"] == "P8-P10"], "Weight_g", "P8–P10 Weight", "Weight (g)")
    st2 = plot_sex_bar_scatter(ax_b2, df_w[df_w["Timepoint_Label"] == "P28"], "Weight_g", "P28 Weight", "Weight (g)")
    st3 = plot_sex_bar_scatter(ax_b3, df_w[df_w["Timepoint_Label"] == "Adult"], "Weight_g", "Adult Weight", "Weight (g)")

    add_subplot_label(ax_wf, "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_wm, "B", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_b1, "C", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_b2, "D", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_b3, "E", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    return fig, st_f + st_m + st1 + st2 + st3


def plot_sex_dvc_hourly_line(ax, df_hourly, sex_code, sex_title):
    """Plot DVC hourly activity line plot for Female or Male matching paper style."""
    df_sex = df_hourly[df_hourly["Sex_Clean"] == sex_code].copy()
    ax.axvspan(0, 6, color='#e6f2ff', alpha=0.4, lw=0)
    ax.axvspan(18, 23, color='#e6f2ff', alpha=0.4, lw=0)

    wt_sub = df_sex[df_sex["Genotype"] == "WT"].sort_values("Hour")
    mut_sub = df_sex[df_sex["Genotype"] == "I80T/+"].sort_values("Hour")

    wt_n  = int(wt_sub["N"].dropna().iloc[0]) if not wt_sub.empty else 0
    mut_n = int(mut_sub["N"].dropna().iloc[0]) if not mut_sub.empty else 0

    wt_color  = SEX_GENO_COLORS[f"WT-{sex_code}"]
    mut_color = SEX_GENO_COLORS[f"I80T/+-{sex_code}"]

    if not wt_sub.empty:
        ax.plot(wt_sub["Hour"], wt_sub["Mean"], color=wt_color, label=f"WT {sex_code} (n={wt_n})", linewidth=1.2)
        ax.fill_between(wt_sub["Hour"], wt_sub["Mean"] - wt_sub["SEM"], wt_sub["Mean"] + wt_sub["SEM"], color=wt_color, alpha=0.2, lw=0)

    if not mut_sub.empty:
        ax.plot(mut_sub["Hour"], mut_sub["Mean"], color=mut_color, label=f"I80T/+ {sex_code} (n={mut_n})", linewidth=1.2)
        ax.fill_between(mut_sub["Hour"], mut_sub["Mean"] - mut_sub["SEM"], mut_sub["Mean"] + mut_sub["SEM"], color=mut_color, alpha=0.2, lw=0)

    ax.set_xlabel("Hour of Day", fontsize=8)
    ax.set_ylabel("Distance Traveled (m)", fontsize=8)
    ax.set_title(f"Circadian Activity ({sex_title})", fontsize=9, fontweight="bold")
    ax.set_xlim(0, 23)
    ax.set_xticks([0, 6, 12, 18, 23])
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(0.8)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.tick_params(direction="out", length=3, width=0.8, labelsize=8)
    ax.legend(frameon=False, fontsize=7, loc="upper right")
    apply_clean_yticks(ax)


def plot_sex_dvc(data):
    """Figure: Circadian Activity (Hourly Line Plots + Total Dark Activity Bar Plot)."""
    print("  → DVC Circadian Activity by Sex")
    df_cage   = data["dvc_cage"].dropna(subset=["Sum_All_Dark", "SexGeno"])
    df_hourly = data["dvc_hourly"]

    fig = plt.figure(figsize=(6.93, 5.5))
    gs  = gridspec.GridSpec(2, 2, height_ratios=[1.1, 1.0], hspace=0.45, wspace=0.35)

    ax_f   = fig.add_subplot(gs[0, 0])
    ax_m   = fig.add_subplot(gs[0, 1])
    ax_tot = fig.add_subplot(gs[1, :])

    # Top panel A: Hourly circadian activity lines
    plot_sex_dvc_hourly_line(ax_f, df_hourly, "F", "Females")
    plot_sex_dvc_hourly_line(ax_m, df_hourly, "M", "Males")

    # Bottom panel: Total Dark Activity
    stats_out = plot_sex_bar_scatter(ax_tot, df_cage, "Sum_All_Dark",
                                     "Total Activity (Dark Phase)",
                                     "Summed Activity (m)")

    add_subplot_label(ax_f, "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_tot, "B", x=-0.06, y=1.08, fontsize=10, fontweight="bold")

    return fig, stats_out


def plot_sex_figure_1_behavior_full(data):
    """Full Figure 1 equivalent (Behavior & Development) broken down by sex x genotype, containing both Line and Bar plots."""
    print("  → Full Figure 1 Behavior by Sex")
    df_w      = data["weights"].dropna(subset=["Weight_g", "SexGeno"])
    df_loc    = data["of_loc"].dropna(subset=["Distance (m)", "SexGeno"])
    df_anx    = data["of_anx"].dropna(subset=["Center_Outer_Time_Ratio", "SexGeno"])
    df_cage   = data["dvc_cage"].dropna(subset=["Sum_All_Dark", "SexGeno"])
    df_hourly = data["dvc_hourly"]
    df_t_ent  = data["tmaze_ent"].dropna(subset=["Distance (m)", "SexGeno"])
    df_t_alt  = data["tmaze_alt"].dropna(subset=["Percent_Alternations", "SexGeno"])

    fig = plt.figure(figsize=(6.93, 11.2))
    gs  = gridspec.GridSpec(5, 1, height_ratios=[1.0, 1.0, 1.0, 1.1, 1.0], hspace=0.48)

    # ROW 1: Developmental Body Weight Line Plots (Female & Male longitudinal lines matching Fig 1B)
    gs_r1 = gridspec.GridSpecFromSubplotSpec(1, 2, subplot_spec=gs[0], wspace=0.35)
    ax_wf = fig.add_subplot(gs_r1[0])
    ax_wm = fig.add_subplot(gs_r1[1])
    plot_sex_longitudinal_weights_single(ax_wf, df_w, "F", "Females")
    plot_sex_longitudinal_weights_single(ax_wm, df_w, "M", "Males")
    add_subplot_label(ax_wf, "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    # ROW 2: Developmental Body Weight Timepoint Bar Plots (P8-P10, P28, Adult)
    gs_r2 = gridspec.GridSpecFromSubplotSpec(1, 3, subplot_spec=gs[1], wspace=0.35)
    ax_w1 = fig.add_subplot(gs_r2[0])
    ax_w2 = fig.add_subplot(gs_r2[1])
    ax_w3 = fig.add_subplot(gs_r2[2])
    plot_sex_bar_scatter(ax_w1, df_w[df_w["Timepoint_Label"] == "P8-P10"], "Weight_g", "P8–P10 Weight", "Weight (g)")
    plot_sex_bar_scatter(ax_w2, df_w[df_w["Timepoint_Label"] == "P28"], "Weight_g", "P28 Weight", "Weight (g)")
    plot_sex_bar_scatter(ax_w3, df_w[df_w["Timepoint_Label"] == "Adult"], "Weight_g", "Adult Weight", "Weight (g)")
    add_subplot_label(ax_w1, "B", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    # ROW 3: Open Field (Locomotion & Anxiety)
    gs_r3 = gridspec.GridSpecFromSubplotSpec(1, 2, subplot_spec=gs[2], wspace=0.35)
    ax_loc = fig.add_subplot(gs_r3[0])
    ax_anx = fig.add_subplot(gs_r3[1])
    plot_sex_bar_scatter(ax_loc, df_loc, "Distance (m)", "Open Field Locomotion", "Total Distance (m)")
    plot_sex_bar_scatter(ax_anx, df_anx, "Center_Outer_Time_Ratio", "Open Field Anxiety", "Center:Outer Time Ratio")
    add_subplot_label(ax_loc, "C", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_anx, "D", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    # ROW 4: Circadian Activity (Hourly Lines + Dark Phase Sum)
    gs_r4 = gridspec.GridSpecFromSubplotSpec(1, 3, subplot_spec=gs[3], wspace=0.35, width_ratios=[1, 1, 1])
    ax_df = fig.add_subplot(gs_r4[0])
    ax_dm = fig.add_subplot(gs_r4[1])
    ax_dt = fig.add_subplot(gs_r4[2])
    plot_sex_dvc_hourly_line(ax_df, df_hourly, "F", "Females")
    plot_sex_dvc_hourly_line(ax_dm, df_hourly, "M", "Males")
    plot_sex_bar_scatter(ax_dt, df_cage, "Sum_All_Dark", "Total Activity (Dark Phase)", "Summed Distance (m)")
    add_subplot_label(ax_df, "E", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_dt, "F", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    # ROW 5: T-Maze (Distance, Arm Entries, % Alternation)
    gs_r5 = gridspec.GridSpecFromSubplotSpec(1, 3, subplot_spec=gs[4], wspace=0.35)
    ax_t1 = fig.add_subplot(gs_r5[0])
    ax_t2 = fig.add_subplot(gs_r5[1])
    ax_t3 = fig.add_subplot(gs_r5[2])
    plot_sex_bar_scatter(ax_t1, df_t_ent, "Distance (m)", "T-Maze Distance", "Distance Traveled (m)")
    plot_sex_bar_scatter(ax_t2, df_t_ent, "Total_Arm_Entries", "T-Maze Arm Entries", "Total Entries")
    plot_sex_bar_scatter(ax_t3, df_t_alt, "Percent_Alternations", "T-Maze Alternations", "Alternation (%)")
    add_subplot_label(ax_t1, "G", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_t2, "H", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(ax_t3, "I", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    return fig, {}


def plot_sex_input_resistance(data):
    """Figure: Somatic Input Resistance by sex × genotype."""
    print("  → Input Resistance by Sex")
    df = data["intrinsic"].dropna(subset=["Input_Resistance_MOhm", "SexGeno"])

    fig, ax = plt.subplots(1, 1, figsize=(3.45, 3.2))
    fig.subplots_adjust(left=0.20, right=0.95, top=0.88, bottom=0.22)

    stats_out = plot_sex_bar_scatter(ax, df, "Input_Resistance_MOhm",
                                     "Somatic Input Resistance",
                                     "Input Resistance (MΩ)")
    return fig, stats_out


def plot_sex_ephys_multi(data):
    """Figure: Somatic and Dendritic Excitability (Input Resistance + FI Midpoint + Plateau Area)."""
    print("  → Ephys Somatic & Dendritic Excitability (Input Resistance, FI Midpoint, Plateau Area)")
    df_rin  = data["intrinsic"].dropna(subset=["Input_Resistance_MOhm", "SexGeno"])
    df_fi   = data["fi_midpoint"].dropna(subset=["FI_Midpoint", "SexGeno"])
    df_plat = data["plateau"]
    df_gab  = df_plat[df_plat["Condition"].str.contains("Gabazine", na=False, case=False)].dropna(subset=["Plateau_Area", "SexGeno"])

    fig, axes = plt.subplots(1, 3, figsize=(6.93, 3.2))
    fig.subplots_adjust(wspace=0.35, left=0.08, right=0.96, top=0.88, bottom=0.22)

    st_rin  = plot_sex_bar_scatter(axes[0], df_rin, "Input_Resistance_MOhm", "Somatic Input Resistance", "Input Resistance (MΩ)")
    st_fi   = plot_sex_bar_scatter(axes[1], df_fi, "FI_Midpoint", "Somatic Excitability", "F-I Midpoint (pA)")
    st_plat = plot_sex_bar_scatter(axes[2], df_gab, "Plateau_Area", "Dendritic Excitability", "Plateau Area (mV·s)")

    add_subplot_label(axes[0], "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(axes[1], "B", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(axes[2], "C", x=-0.12, y=1.08, fontsize=10, fontweight="bold")

    return fig, {"rin": st_rin, "fi_midpoint": st_fi, "plateau": st_plat}


def plot_sex_fi_midpoint(data):
    """Figure: F-I Midpoint by sex × genotype."""
    print("  → FI Midpoint by Sex")
    df = data["fi_midpoint"].dropna(subset=["FI_Midpoint", "SexGeno"])

    fig, ax = plt.subplots(1, 1, figsize=(3.45, 3.2))
    fig.subplots_adjust(left=0.20, right=0.95, top=0.88, bottom=0.22)

    stats_out = plot_sex_bar_scatter(ax, df, "FI_Midpoint",
                                     "F-I Curve Midpoint",
                                     "F-I Midpoint (pA)")
    return fig, stats_out


_EI_ISIS = [300, 100, 50, 25, 10]

def plot_sex_ei_line(ax, df_ei, pathway, title, ylabel="E/I Imbalance Index", add_legend=False):
    MARKER_MAP = {"WT-M": "^", "WT-F": "o", "I80T/+-M": "^", "I80T/+-F": "o"}
    df_pw = df_ei[df_ei["Pathway"] == pathway].copy()
    stats_out = []

    for group in SEX_GENO_ORDER:
        means, sems = [], []
        for isi in _EI_ISIS:
            sub = df_pw[(df_pw["SexGeno"] == group) & (df_pw["ISI"] == isi)]["E_I_Imbalance"].dropna()
            means.append(sub.mean() if len(sub) > 0 else np.nan)
            sems.append(sub.sem()  if len(sub) > 1 else 0.0)
        x      = np.arange(len(_EI_ISIS))
        color  = SEX_GENO_COLORS.get(group, "gray")
        ls     = SEX_GENO_LINES.get(group, "-")
        marker = MARKER_MAP.get(group, "o")
        n_vals = [len(df_pw[(df_pw["SexGeno"] == group) & (df_pw["ISI"] == isi)]["E_I_Imbalance"].dropna()) for isi in _EI_ISIS]
        n_str  = f"{min(n_vals)}–{max(n_vals)}" if min(n_vals) != max(n_vals) else str(min(n_vals))
        ax.errorbar(x, means, yerr=sems, marker=marker, markersize=3.5, linewidth=1.0,
                    linestyle=ls, capsize=1.5, capthick=0.8, elinewidth=0.8, color=color, label=f"{group} (n={n_str})")

    for isi in _EI_ISIS:
        for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
            geno_a, geno_b = comp.split(" vs ")
            a = df_pw[(df_pw["SexGeno"] == geno_a) & (df_pw["ISI"] == isi)]["E_I_Imbalance"].dropna().values
            b = df_pw[(df_pw["SexGeno"] == geno_b) & (df_pw["ISI"] == isi)]["E_I_Imbalance"].dropna().values
            u, p = mann_whitney(a, b)
            stats_out.append(dict(Pathway=pathway, ISI=isi, Comparison=comp, SexLabel=sex_label, U=u, P=p, Sig=p_to_sig(p)))

    ax.set_xticks(np.arange(len(_EI_ISIS)))
    ax.set_xticklabels([str(v) for v in _EI_ISIS], fontsize=8)
    ax.set_xlabel("ISI (ms)", fontsize=8)
    ax.set_ylabel(ylabel, fontsize=8)
    ax.set_title(title, fontsize=9, fontweight="bold")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(0.8)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.tick_params(direction="out", length=3, width=0.8, labelsize=8)
    ax.set_ylim(0.2, 1.0)
    apply_clean_yticks(ax)

    if add_legend:
        ax.legend(frameon=False, fontsize=6, loc="lower right")

    return stats_out


def plot_sex_ei_imbalance(data):
    print("  → E/I Imbalance by Sex (all pathways, all ISIs)")
    df_ei = data["ei"]
    pathways_info = [("Perforant", "ECIII (Perforant)"), ("Schaffer", "CA3 Apical (Schaffer)"), ("Basal_Stratum_Oriens", "CA3 Basal")]
    fig, axes = plt.subplots(1, 3, figsize=(6.93, 2.8))
    fig.subplots_adjust(wspace=0.35, left=0.09, right=0.96, top=0.85, bottom=0.22)
    all_stats = []
    for col_idx, (ax, (pw, pw_label)) in enumerate(zip(axes, pathways_info)):
        s = plot_sex_ei_line(ax, df_ei, pw, pw_label, ylabel="E/I Imbalance Index" if col_idx == 0 else "", add_legend=(col_idx == 2))
        all_stats.extend(s)
        add_subplot_label(ax, chr(65 + col_idx), x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    return fig, all_stats


def plot_sex_plateau(data):
    print("  → Plateau Area by Sex")
    df = data["plateau"]
    df_gab = df[df["Condition"].str.contains("Gabazine", na=False, case=False)].copy()
    df_gab = df_gab.dropna(subset=["Plateau_Area", "SexGeno"])
    fig, ax = plt.subplots(1, 1, figsize=(3.45, 3.2))
    fig.subplots_adjust(left=0.20, right=0.95, top=0.88, bottom=0.22)
    stats_out = plot_sex_bar_scatter(ax, df_gab, "Plateau_Area", "Plateau Area (Gabazine)", "Plateau Area (mV·s)")
    return fig, stats_out


def plot_sex_open_field(data):
    print("  → Open Field by Sex")
    df_loc = data["of_loc"].dropna(subset=["Distance (m)", "SexGeno"])
    df_anx = data["of_anx"].dropna(subset=["Center_Outer_Time_Ratio", "SexGeno"])
    fig, axes = plt.subplots(1, 2, figsize=(6.93, 3.2))
    fig.subplots_adjust(wspace=0.35, left=0.10, right=0.96, top=0.88, bottom=0.22)
    stats_loc = plot_sex_bar_scatter(axes[0], df_loc, "Distance (m)", "Open Field Locomotion", "Distance (m)")
    stats_anx = plot_sex_bar_scatter(axes[1], df_anx, "Center_Outer_Time_Ratio", "Open Field Anxiety", "Center:Outer Time Ratio")
    add_subplot_label(axes[0], "A", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    add_subplot_label(axes[1], "B", x=-0.12, y=1.08, fontsize=10, fontweight="bold")
    return fig, {"location": stats_loc, "anxiety": stats_anx}


def plot_sex_olm(data):
    print("  → OLM by Sex")
    df = data["olm"].dropna(subset=["Testing_DI", "SexGeno"])
    fig, ax = plt.subplots(1, 1, figsize=(3.45, 3.2))
    fig.subplots_adjust(left=0.20, right=0.95, top=0.88, bottom=0.22)
    stats_out = plot_sex_bar_scatter(ax, df, "Testing_DI", "OLM – Testing Discrimination Index", "Discrimination Index")
    return fig, stats_out


# ─────────────────────────────────────────────────────────────────────────────
# STATS COLLECTION → Master_Stats_Summary tabs
# ─────────────────────────────────────────────────────────────────────────────

def collect_all_stats(data):
    """Compile a single DataFrame of all sex-stratified stats."""
    all_rows = []

    # ── 1. Body Weights ──────────────────────────────────────────────────────
    df_w = data["weights"].dropna(subset=["Weight_g", "SexGeno"])
    for tp in ["P8-P10", "P28", "Adult"]:
        sub = df_w[df_w["Timepoint_Label"] == tp]
        for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
            all_rows.append(build_stats_row(
                "Supplemental – Sex Stratified", f"Body Weight – {tp}",
                f"Body Weight ({tp}) (g)", "N/A", tp,
                sub, "Weight_g", sex_label, comp
            ))

    # ── 2. DVC ──────────────────────────────────────────────────────────────
    df_dvc = data["dvc_cage"].dropna(subset=["Sum_All_Dark", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "Circadian Activity",
            "DVC Summed Dark Activity (m)", "N/A", "All Dark Hours",
            df_dvc, "Sum_All_Dark", sex_label, comp
        ))

    # ── 3. FI Midpoint ──────────────────────────────────────────────────────
    df_fi = data["fi_midpoint"].dropna(subset=["FI_Midpoint", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "Somatic Excitability",
            "F-I Curve Midpoint (pA)", "N/A", "N/A",
            df_fi, "FI_Midpoint", sex_label, comp
        ))

    # ── 4. E/I Imbalance (all ISIs, all 3 pathways) ──────────────────────────
    df_ei = data["ei"].dropna(subset=["E_I_Imbalance", "SexGeno"])
    for pw in ["Perforant", "Schaffer", "Basal_Stratum_Oriens"]:
        df_pw = df_ei[df_ei["Pathway"] == pw]
        for isi in _EI_ISIS:
            df_isi = df_pw[df_pw["ISI"] == isi]
            for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
                geno_a, geno_b = comp.split(" vs ")
                grp_a = df_isi[df_isi["SexGeno"] == geno_a]["E_I_Imbalance"].dropna()
                grp_b = df_isi[df_isi["SexGeno"] == geno_b]["E_I_Imbalance"].dropna()
                u, p  = mann_whitney(grp_a.values, grp_b.values)
                def _f(x): return round(float(x), 4) if pd.notna(x) and not np.isnan(float(x)) else np.nan
                all_rows.append(dict(
                    Figure="Supplemental – Sex Stratified",
                    Subpanel=f"E/I Imbalance – {pw}",
                    Metric=f"E/I Imbalance Index [{sex_label}]",
                    Pathway=pw,
                    Condition=f"ISI {isi} ms",
                    WT_Mean=_f(grp_a.mean()), WT_SEM=_f(grp_a.sem()), WT_N=len(grp_a),
                    I80T_Mean=_f(grp_b.mean()), I80T_SEM=_f(grp_b.sem()), I80T_N=len(grp_b),
                    Test_Used="Mann-Whitney U (sex-stratified, per ISI)",
                    Statistic=_f(u), P_Value=_f(p), Significance=p_to_sig(p),
                    Notes=f"Sex-stratified: {comp} | ISI {isi} ms",
                ))

    # ── 5. Plateau Area ─────────────────────────────────────────────────────
    df_plat = data["plateau"]
    df_gab  = df_plat[df_plat["Condition"].str.contains("Gabazine", na=False, case=False)].dropna(subset=["Plateau_Area", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "Dendritic Excitability",
            "Plateau Area (mV·s)", "Both Pathways", "Gabazine condition",
            df_gab, "Plateau_Area", sex_label, comp
        ))

    # ── 6a. Open Field Locomotion ─────────────────────────────────────────
    df_loc = data["of_loc"].dropna(subset=["Distance (m)", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "Open Field Locomotion",
            "Open Field Distance (m)", "N/A", "Habituation Day 1 Trial 1",
            df_loc, "Distance (m)", sex_label, comp
        ))

    # ── 6b. Open Field Anxiety ────────────────────────────────────────────
    df_anx = data["of_anx"].dropna(subset=["Center_Outer_Time_Ratio", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "Open Field Anxiety",
            "Open Field Center:Outer Time Ratio", "N/A", "N/A",
            df_anx, "Center_Outer_Time_Ratio", sex_label, comp
        ))

    # ── 7. T-Maze ─────────────────────────────────────────────────────────
    df_t_ent = data["tmaze_ent"].dropna(subset=["Distance (m)", "Total_Arm_Entries", "SexGeno"])
    df_t_alt = data["tmaze_alt"].dropna(subset=["Percent_Alternations", "SexGeno"])

    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "T-Maze Distance",
            "T-Maze Distance Traveled (m)", "N/A", "Spontaneous Alternation Trial",
            df_t_ent, "Distance (m)", sex_label, comp
        ))
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "T-Maze Entries",
            "T-Maze Total Arm Entries", "N/A", "Spontaneous Alternation Trial",
            df_t_ent, "Total_Arm_Entries", sex_label, comp
        ))
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "T-Maze Alternations",
            "T-Maze Percent Alternations (%)", "N/A", "Spontaneous Alternation Trial",
            df_t_alt, "Percent_Alternations", sex_label, comp
        ))

    # ── 8. OLM ────────────────────────────────────────────────────────────
    df_olm = data["olm"].dropna(subset=["Testing_DI", "SexGeno"])
    for comp, sex_label in [("WT-M vs I80T/+-M", "Male"), ("WT-F vs I80T/+-F", "Female")]:
        all_rows.append(build_stats_row(
            "Supplemental – Sex Stratified", "OLM",
            "OLM Testing Discrimination Index", "N/A", "Testing Stage",
            df_olm, "Testing_DI", sex_label, comp
        ))

    return pd.DataFrame(all_rows)


COLUMNS = [
    "Figure", "Subpanel", "Metric", "Pathway", "Condition",
    "WT_Mean", "WT_SEM", "WT_N",
    "I80T_Mean", "I80T_SEM", "I80T_N",
    "Test_Used", "Statistic", "P_Value", "Significance",
    "Notes",
]

def export_stats(df_sex_stats):
    """Append sex-stratified stats to Master_Stats_Summary CSV and add tabs to XLSX."""
    if os.path.exists(MASTER_CSV):
        df_existing = pd.read_csv(MASTER_CSV)
        mask = df_existing.get("Figure", pd.Series(dtype=str)).str.startswith("Supplemental – Sex", na=False)
        df_existing = df_existing[~mask]
        df_combined = pd.concat([df_existing, df_sex_stats], ignore_index=True)
    else:
        df_combined = df_sex_stats

    df_combined.to_csv(MASTER_CSV, index=False)
    print(f"  ✓ Updated CSV  → {MASTER_CSV} ({len(df_sex_stats)} sex-stratified rows appended)")

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font

        if os.path.exists(MASTER_XLSX):
            wb = load_workbook(MASTER_XLSX)
        else:
            import openpyxl
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

        stale = [s for s in wb.sheetnames if "Sex" in s]
        for s in stale:
            del wb[s]

        ws_all = wb.create_sheet("Sex Stratified – All")
        ws_all.append(COLUMNS)
        for r in df_sex_stats[COLUMNS].itertuples(index=False):
            ws_all.append(list(r))

        metric_groups = {
            "Sex - Weights":      lambda d: d["Subpanel"].str.startswith("Body Weight", na=False),
            "Sex - DVC":          lambda d: d["Subpanel"] == "Circadian Activity",
            "Sex - FI Midpoint":  lambda d: d["Subpanel"] == "Somatic Excitability",
            "Sex - EI Imbalance": lambda d: d["Subpanel"].str.startswith("E/I", na=False),
            "Sex - Plateau Area": lambda d: d["Subpanel"] == "Dendritic Excitability",
            "Sex - Open Field":   lambda d: d["Subpanel"].str.startswith("Open Field", na=False),
            "Sex - T-Maze":       lambda d: d["Subpanel"].str.startswith("T-Maze", na=False),
            "Sex - OLM":          lambda d: d["Subpanel"] == "OLM",
        }
        for sheet_name, mask_fn in metric_groups.items():
            subset = df_sex_stats[mask_fn(df_sex_stats)]
            if subset.empty:
                continue
            ws = wb.create_sheet(sheet_name[:31])
            ws.append(COLUMNS)
            for r in subset[COLUMNS].itertuples(index=False):
                ws.append(list(r))
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col_cells in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)

        for cell in ws_all[1]:
            cell.font = Font(bold=True)
        for col_cells in ws_all.columns:
            max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
            ws_all.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)

        wb.save(MASTER_XLSX)
        print(f"  ✓ Updated XLSX → {MASTER_XLSX}")

    except ImportError:
        print("  ⚠ openpyxl not installed – skipping XLSX export.")


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main():
    setup_publication_style()

    print("\n=== Sex-Stratified Figure Generation ===\n")
    print("Loading data …")
    data = load_all_data()

    # ── 1. Full Figure 1 Behavior equivalent by Sex ──────────────────────────
    print("\n[1] Full Figure 1 Behavior (Body Weight, Open Field, Circadian, T-Maze)")
    fig, _ = plot_sex_figure_1_behavior_full(data)
    save_fig_png_svg(fig, "SexStratified_Figure1_Behavior")

    # ── 2. Body Weights ──────────────────────────────────────────────────────
    print("\n[2] Body Weights (Combined, Line Plots, Bar Plots)")
    fig, _ = plot_sex_weights(data)
    save_fig_png_svg(fig, "SexStratified_Weights")

    fig_lp, _ = plot_sex_weights_line_plots(data)
    save_fig_png_svg(fig_lp, "SexStratified_Weights_LinePlots")

    fig_bp, _ = plot_sex_weights_bar_plots(data)
    save_fig_png_svg(fig_bp, "SexStratified_Weights_BarPlots")

    # ── 3. DVC Circadian ─────────────────────────────────────────────────────
    print("\n[3] DVC – Circadian Activity")
    fig, _ = plot_sex_dvc(data)
    save_fig_png_svg(fig, "SexStratified_DVC_Circadian_Activity")

    # ── 4. Ephys Multi (Input Resistance, FI Midpoint, Plateau Area) ─────────
    print("\n[4] Ephys Somatic & Dendritic Excitability")
    fig, _ = plot_sex_ephys_multi(data)
    save_fig_png_svg(fig, "SexStratified_Ephys_Figure")

    # ── 4b. Input Resistance (Single Panel) ──────────────────────────────────
    print("\n[4b] Input Resistance")
    fig, _ = plot_sex_input_resistance(data)
    save_fig_png_svg(fig, "SexStratified_Input_Resistance")

    # ── 5. FI Midpoint (Single Panel) ────────────────────────────────────────
    print("\n[5] FI Midpoint")
    fig, _ = plot_sex_fi_midpoint(data)
    save_fig_png_svg(fig, "SexStratified_FI_Midpoint")

    # ── 6. Plateau Area (Single Panel) ───────────────────────────────────────
    print("\n[6] Plateau Area")
    fig, _ = plot_sex_plateau(data)
    save_fig_png_svg(fig, "SexStratified_Plateau_Area")

    # ── 7. Open Field (Single Panel Pair) ────────────────────────────────────
    print("\n[7] Open Field Location + Anxiety")
    fig, _ = plot_sex_open_field(data)
    save_fig_png_svg(fig, "SexStratified_Open_Field")

    # ── 8. E/I Imbalance ─────────────────────────────────────────────────────
    print("\n[8] E/I Imbalance")
    fig, _ = plot_sex_ei_imbalance(data)
    save_fig_png_svg(fig, "SexStratified_EI_Imbalance")

    # ── 9. OLM ───────────────────────────────────────────────────────────────
    print("\n[9] OLM – Testing Discrimination Index")
    fig, _ = plot_sex_olm(data)
    save_fig_png_svg(fig, "SexStratified_OLM")

    # ── Stats export ──────────────────────────────────────────────────────────
    print("\n[→] Compiling stats …")
    df_sex_stats = collect_all_stats(data)
    for c in COLUMNS:
        if c not in df_sex_stats.columns:
            df_sex_stats[c] = np.nan
    df_sex_stats = df_sex_stats[COLUMNS]

    print(df_sex_stats[["Subpanel", "Metric", "Pathway", "Notes",
                         "WT_Mean", "I80T_Mean", "P_Value", "Significance"]].to_string())

    print("\n[→] Exporting stats to Master_Stats_Summary …")
    export_stats(df_sex_stats)

    print("\n✓ Done. All sex-stratified figures saved to paper_figures/")


if __name__ == "__main__":
    main()
